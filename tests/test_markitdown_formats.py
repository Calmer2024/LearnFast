from __future__ import annotations

import csv
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps" / "api"))

from learnfast.services.converter import convert_to_markdown


FIXTURE_DIR = ROOT / ".learnfast-data" / "test-fixtures" / "markitdown"


def write_text_fixtures() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    md = FIXTURE_DIR / "sample.md"
    txt = FIXTURE_DIR / "sample.txt"
    csv_path = FIXTURE_DIR / "sample.csv"

    md.write_text("# Markdown Fixture\n\nLearnFast markdown body.", encoding="utf-8")
    txt.write_text("Plain text fixture for LearnFast.", encoding="utf-8")
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["concept", "score"])
        writer.writerow(["retrieval", "95"])

    return {"md": md, "txt": txt, "csv": csv_path}


def write_xlsx_fixture() -> Path:
    from openpyxl import Workbook

    path = FIXTURE_DIR / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scores"
    sheet.append(["concept", "score"])
    sheet.append(["chunking", 88])
    workbook.save(path)
    return path


def write_docx_fixture() -> Path:
    from docx import Document

    path = FIXTURE_DIR / "sample.docx"
    document = Document()
    document.add_heading("DOCX Fixture", level=1)
    document.add_paragraph("LearnFast extracts Word documents.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "concept"
    table.cell(0, 1).text = "status"
    table.cell(1, 0).text = "memory"
    table.cell(1, 1).text = "ready"
    document.save(path)
    return path


def write_pptx_fixture() -> Path:
    from pptx import Presentation

    path = FIXTURE_DIR / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPTX Fixture"
    slide.placeholders[1].text = "LearnFast extracts slide text."
    presentation.save(path)
    return path


def write_pdf_fixture() -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = FIXTURE_DIR / "sample.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.drawString(72, 720, "PDF Fixture")
    pdf.drawString(72, 700, "LearnFast extracts PDF text.")
    pdf.save()
    return path


def write_epub_fixture() -> Path:
    path = FIXTURE_DIR / "sample.epub"
    mimetype = "application/epub+zip"
    container_xml = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""
    content_opf = """<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>EPub Fixture</dc:title>
    <dc:language>en</dc:language>
    <dc:identifier id="bookid">learnfast-epub-fixture</dc:identifier>
  </metadata>
  <manifest>
    <item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine>
    <itemref idref="chapter"/>
  </spine>
</package>
"""
    chapter = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml">
  <head><title>EPub Fixture</title></head>
  <body>
    <h1>EPub Fixture</h1>
    <p>LearnFast extracts ebook text.</p>
  </body>
</html>
"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container_xml)
        archive.writestr("OEBPS/content.opf", content_opf)
        archive.writestr("OEBPS/chapter.xhtml", chapter)
    return path


def build_fixtures() -> dict[str, Path]:
    fixtures = write_text_fixtures()
    fixtures.update(
        {
            "xlsx": write_xlsx_fixture(),
            "docx": write_docx_fixture(),
            "pptx": write_pptx_fixture(),
            "pdf": write_pdf_fixture(),
            "epub": write_epub_fixture(),
        }
    )
    return fixtures


def main() -> int:
    expected = {
        "md": "Markdown Fixture",
        "txt": "Plain text fixture",
        "csv": "retrieval",
        "xlsx": "chunking",
        "docx": "DOCX Fixture",
        "pptx": "PPTX Fixture",
        "pdf": "PDF Fixture",
        "epub": "EPub Fixture",
    }
    failures: list[str] = []

    for label, path in build_fixtures().items():
        try:
            markdown = convert_to_markdown(path)
            ok = expected[label] in markdown
            status = "PASS" if ok else "FAIL"
            print(f"{status} {label.upper():5} {path.name} -> {len(markdown)} chars")
            if not ok:
                failures.append(f"{label}: expected marker {expected[label]!r} not found")
        except Exception as exc:
            print(f"FAIL {label.upper():5} {path.name} -> {exc}")
            failures.append(f"{label}: {exc}")

    if failures:
        print("\nFailures:")
        for failure in failures:
            print(f"- {failure}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
